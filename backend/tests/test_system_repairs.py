"""Regression coverage for upload failures, provenance and proof integrity."""
import uuid
from pathlib import Path

import pytest
from sqlmodel import select

from app.core.proof_sandbox import solve
from app.models import Audit, Document, Passport


def test_boolean_contradiction_and_sum():
    assert solve(["approved"], ["not approved"], [], 2)["satisfiable"] is False
    assert solve(["a + b + c == 100"], [], [
        {"variable": "a", "operator": "eq", "value": 40},
        {"variable": "b", "operator": "eq", "value": 40},
        {"variable": "c", "operator": "eq", "value": 30},
    ], 2)["satisfiable"] is False


@pytest.mark.parametrize("expression", ["__import__('os').system('whoami')", "Currency is INR", "x / 0 == 1"])
def test_proof_rejects_unsupported_expressions(expression):
    with pytest.raises(ValueError):
        solve([expression], [], [], 2)


def test_missing_checkpoint_is_unavailable(tmp_path):
    from app.core.rope_inference import load_rope_model
    with pytest.raises(RuntimeError, match="missing"):
        load_rope_model(str(tmp_path / "missing.pt"))


def test_docx_block_order(tmp_path):
    import docx
    from app.core.parsers import parse_file
    file = tmp_path / "ordered.docx"
    document = docx.Document()
    document.add_paragraph("Before the table.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Inside the table."
    document.add_paragraph("After the table.")
    document.save(file)
    text = parse_file(file).raw_text
    assert text.index("Before") < text.index("Inside") < text.index("After")


def test_protected_operations_require_login(client):
    for path, body in [("/api/v1/rope/train", {}), ("/api/v1/proof/solve", {}), ("/api/v1/inject/candidate", {"original_text": "a", "injected_text": "b"})]:
        assert client.post(path, json=body).status_code == 401
    assert client.get("/api/v1/metrics").status_code == 401


def test_no_file_does_not_create_audit(client, db, superuser_token_headers):
    before = len(db.exec(select(Audit)).all())
    response = client.post("/api/v1/audits", data={"title": "Empty"}, headers=superuser_token_headers)
    assert response.status_code == 400
    assert len(db.exec(select(Audit)).all()) == before


def test_failed_batch_rolls_back_and_removes_files(client, db, superuser_token_headers, monkeypatch):
    from app.api.routes import audits
    from app.core.jobs import StageError
    from app.core.config import settings
    original = audits.ingest_document
    def ingest(session, audit, **kwargs):
        if kwargs["filename"] == "bad.txt":
            raise StageError("ingest", "deliberately rejected")
        return original(session, audit, **kwargs)
    monkeypatch.setattr(audits, "ingest_document", ingest)
    audit_before = len(db.exec(select(Audit)).all())
    doc_before = len(db.exec(select(Document)).all())
    stored_before = set(Path(settings.STORAGE_DIR).rglob("*"))
    response = client.post("/api/v1/audits", data={"title": "Atomic batch"}, files=[
        ("files", ("good.txt", b"Readable first document.", "text/plain")),
        ("files", ("bad.txt", b"Rejected second document.", "text/plain")),
    ], headers=superuser_token_headers)
    assert response.status_code == 422, response.text
    assert len(db.exec(select(Audit)).all()) == audit_before
    assert len(db.exec(select(Document)).all()) == doc_before
    assert not [p for p in set(Path(settings.STORAGE_DIR).rglob("*")) - stored_before if p.is_file()]


def test_metrics_and_empty_evaluations(client, superuser_token_headers):
    assert client.get("/api/v1/rope/metrics", headers=superuser_token_headers).status_code == 200
    response = client.get("/api/v1/challenges/results", headers=superuser_token_headers)
    assert response.status_code == 200
    assert response.json()["results"] == []
    assert response.json()["quality_report"]["sufficient_trials"] is False


def test_invalid_evidence_is_client_error(client, superuser_token_headers):
    response = client.post("/api/v1/judge/adversarial", json={"claim_text": "Example claim", "evidence_list": [{"id": "invalid", "quote": "Example evidence"}]}, headers=superuser_token_headers)
    assert response.status_code == 422


def test_stored_score_breakdown_and_passport_reissue(client, db, superuser_token_headers):
    from app.core.pipeline import issue_passport
    from app.core.security import ensure_keypair, verify_signature
    from app.core.config import settings
    from app.models import User
    from app.core.pipeline import ingest_document, save_upload
    user = db.exec(select(User).where(User.is_superuser.is_(True))).first()
    audit = Audit(title="Independent passport regression", owner_id=user.id, status="completed", ai_score=90, reviewed_score=90)
    db.add(audit)
    db.commit()
    ingest_document(db, audit, filename="passport.txt", path=save_upload(audit.id, "passport.txt", b"A real source document."))
    from app.core.canonical import source_set_hash
    audit.source_set_hash = source_set_hash([doc.text_hash for doc in db.exec(select(Document).where(Document.audit_id == audit.id)).all()])
    db.add(audit)
    db.commit()
    passport = issue_passport(db, audit)
    response = client.get(f"/api/v1/audits/{audit.id}/score-breakdown", headers=superuser_token_headers)
    assert response.status_code == 200, response.text
    token = passport.verify_token
    revision = passport.revision
    refreshed = issue_passport(db, audit)
    assert refreshed.verify_token == token
    assert refreshed.revision == revision + 1
    assert refreshed.trust_score == audit.reviewed_score
    _, public = ensure_keypair(settings.ECDSA_PRIVATE_KEY_PATH, settings.ECDSA_PUBLIC_KEY_PATH)
    from app.core.passport_payload import signed_payload
    assert refreshed.signature.startswith("v2:")
    assert verify_signature(public, signed_payload(refreshed), refreshed.signature[3:])
    refreshed.trust_score = 99.0
    assert not verify_signature(public, signed_payload(refreshed), refreshed.signature[3:])
    db.rollback()


def test_policy_version_and_missing_data(db):
    import json
    from app.core.policy_engine import create_policy, update_policy_version, evaluate_policy
    policy = create_policy(db, name="regression-" + uuid.uuid4().hex, rules_json=json.dumps({"conditions": [{"variable": "missing_amount", "operator": "lte", "value": 10}]}))
    result = evaluate_policy(db, str(policy.id), claim={"text": "No amount provided"})
    assert result.overall_state == "uncertain"
    assert result.policy_version == "1"
    policy = update_policy_version(db, policy, new_enabled=False)
    assert json.loads(policy.rules_json)["version"] == 2
    assert not policy.is_active
    assert evaluate_policy(db, str(policy.id), claim={}).overall_state == "not_applicable"


def test_static_api_misses_are_not_html(client):
    response = client.get("/api/v1/not-a-real-route", headers={"Accept": "text/html"})
    assert response.status_code == 404
    assert "text/html" not in response.headers.get("content-type", "")


def test_model_status_does_not_fabricate_health_or_latency(client, superuser_token_headers):
    assert client.get("/api/v1/utils/models-status/").status_code == 401
    response = client.get("/api/v1/utils/models-status/", headers=superuser_token_headers)
    assert response.status_code == 200
    data = response.json()
    assert data["availability_measured"] is False
    assert data["active_primary_llm"] is None
    assert all(model["latency"] == "Not measured" for model in data["models"])
    assert all("ONLINE" not in model["status"] for model in data["models"])


def test_concurrent_verifier_requests_share_one_loaded_model(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    import time
    from types import SimpleNamespace
    from app.core import rope_inference as inference
    checkpoint = tmp_path / "model.pt"
    checkpoint.write_bytes(b"fixture")
    constructed = []
    class Model:
        def __init__(self, **kwargs):
            constructed.append(self)
            time.sleep(0.02)
        def to(self, device): return self
        def load_state_dict(self, state): pass
        def eval(self): pass
    monkeypatch.setattr(inference, "TathyaRoPEVerifier", Model)
    monkeypatch.setattr(inference, "DocumentTokenizer", lambda **kwargs: SimpleNamespace(vocab_size=10))
    monkeypatch.setattr(inference.torch, "load", lambda *args, **kwargs: {"model_state_dict": {}})
    monkeypatch.setattr(inference.torch.cuda, "is_available", lambda: False)
    inference.invalidate_rope_model()
    try:
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(lambda _: inference.load_rope_model(str(checkpoint)), range(8)))
        assert len(constructed) == 1
        assert all(item[0] is constructed[0] and item[1] is results[0][1] for item in results)
    finally:
        inference.invalidate_rope_model()
