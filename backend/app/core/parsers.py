"""Parser foundation (Phase 1 + Phase 2).

Locked mapping:
    PDF  -> pdfplumber          (do NOT switch to PyMuPDF)
    DOCX -> python-docx
    XLSX -> openpyxl
    TXT/MD -> direct UTF-8 read
    Other -> MarkItDown fallback (declared offset limitations)

Every parser returns the same representation (ParsedDocument) and records
which parser produced the result (`parser` field: honest provenance).

`locations` are ordered char-span -> location-contract entries so any
offset in raw_text can be resolved to a PDF page+bbox / DOCX paragraph /
XLSX cell without trusting an LLM. Coordinates are never fabricated:
when a format cannot provide precise geometry the location states so.

Failure strategy (locked):
    primary parser fails -> MarkItDown fallback -> ParserError (no loops)
"""

from __future__ import annotations

import logging
import mimetypes
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.core.canonical import (
    location_docx,
    location_pdf,
    location_text,
    location_xlsx,
)

PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv"}

# Safety cap: refuse absurdly large extractions instead of exhausting memory.
MAX_DOCUMENT_CHARS = 5_000_000

PARSER_PDFPLUMBER = "pdfplumber"
PARSER_PYTHON_DOCX = "python-docx"
PARSER_OPENPYXL = "openpyxl"
PARSER_DIRECT = "direct-utf8"
PARSER_MARKITDOWN = "markitdown"


@dataclass
class CharLocation:
    start: int
    end: int
    location: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {"start": self.start, "end": self.end, "location": self.location}


@dataclass
class ParsedDocument:
    raw_text: str
    format: str  # pdf | docx | xlsx | text | fallback
    mime_type: str
    metadata: dict[str, Any] = field(default_factory=dict)
    locations: list[CharLocation] = field(default_factory=list)
    parser: str = PARSER_DIRECT  # which parser produced raw_text
    fallback_used: bool = False  # True when MarkItDown produced the text

    def location_at(self, char_start: int) -> dict[str, Any]:
        for entry in self.locations:
            if entry.start <= char_start < entry.end:
                return entry.location
        return location_text(char_start, char_start)

    def location_span(self, char_start: int, char_end: int) -> dict[str, Any]:
        """Best location for a raw [start, end) span (start-anchored)."""
        return self.location_at(char_start)

    def locations_json(self) -> list[dict[str, Any]]:
        return [e.to_dict() for e in self.locations]


class ParserError(ValueError):
    pass


def _guess_mime(path: Path) -> str:
    mime, _ = mimetypes.guess_type(path.name)
    return mime or "application/octet-stream"


# ---------------------------------------------------------------------------
# PDF -> pdfplumber (real word bboxes, never fabricated)
# ---------------------------------------------------------------------------


def _parse_pdf(path: Path) -> ParsedDocument:
    import pdfplumber

    # pdfminer logs crop-box warnings for many real-world PDFs; keep logs clean
    logging.getLogger("pdfminer").setLevel(logging.ERROR)

    parts: list[str] = []
    locations: list[CharLocation] = []
    cursor = 0
    page_count = 0
    empty_pages = 0
    bbox_available = False

    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        for page_no, page in enumerate(pdf.pages, start=1):
            try:
                words = page.extract_words()
            except Exception:  # noqa: BLE001 - single page may be broken
                words = []
            if not words:
                empty_pages += 1
                continue

            # group words into visual lines (tolerance in points)
            lines: list[list[dict[str, Any]]] = []
            current_line: list[dict[str, Any]] = []
            current_top: float | None = None
            for word in words:
                top = float(word["top"])
                if current_top is None or abs(top - current_top) <= 2.0:
                    current_line.append(word)
                    if current_top is None:
                        current_top = top
                else:
                    lines.append(current_line)
                    current_line = [word]
                    current_top = top
            if current_line:
                lines.append(current_line)

            for line in lines:
                line.sort(key=lambda w: float(w["x0"]))
                start = cursor
                pieces: list[str] = []
                spans: list[tuple[int, int, dict[str, Any]]] = []
                pos = 0
                for index, word in enumerate(line):
                    if index:
                        pieces.append(" ")
                        pos += 1
                    text = str(word["text"])
                    word_start = pos
                    pieces.append(text)
                    pos += len(text)
                    spans.append((word_start, pos, word))
                line_text = "".join(pieces)
                parts.append(line_text)
                for word_start, word_end, word in spans:
                    locations.append(
                        CharLocation(
                            start=start + word_start,
                            end=start + word_end,
                            location=location_pdf(
                                page_no,
                                [
                                    round(float(word["x0"]), 2),
                                    round(float(word["top"]), 2),
                                    round(float(word["x1"]), 2),
                                    round(float(word["bottom"]), 2),
                                ],
                            ),
                        )
                    )
                    bbox_available = True
                cursor += len(line_text)
                parts.append("\n")
                cursor += 1
            parts.append("\n")
            cursor += 1

    raw = "".join(parts).strip("\n")
    return ParsedDocument(
        raw_text=raw,
        format="pdf",
        mime_type="application/pdf",
        parser=PARSER_PDFPLUMBER,
        metadata={
            "page_count": page_count,
            "empty_pages": empty_pages,
            "bbox_available": bbox_available,
        },
        locations=_trim(locations, len(raw)),
    )


# ---------------------------------------------------------------------------
# DOCX -> python-docx (paragraphs + tables, paragraph-indexed locations)
# ---------------------------------------------------------------------------

_HEADING_PREFIX = re.compile(r"^Heading\s*\d*$", re.IGNORECASE)


def _parse_docx(path: Path) -> ParsedDocument:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    locations: list[CharLocation] = []
    cursor = 0
    para_index = 0
    table_count = 0
    headings: list[str] = []

    def _emit_line(
        text: str, loc: dict[str, Any] | None, cell_spans: list[tuple[int, int, dict[str, Any]]] | None = None
    ) -> None:
        nonlocal cursor
        start = cursor
        parts.append(text)
        cursor += len(text)
        if cell_spans:
            for span_start, span_end, cell_loc in cell_spans:
                locations.append(
                    CharLocation(
                        start=start + span_start,
                        end=start + span_end,
                        location=cell_loc,
                    )
                )
        elif loc is not None:
            locations.append(
                CharLocation(start=start, end=cursor, location=loc)
            )
        parts.append("\n\n")
        cursor += 2

    for para in document.paragraphs:
        text = para.text
        style_name = (para.style.name if para.style is not None else "") or ""
        if _HEADING_PREFIX.match(style_name):
            level = re.search(r"\d+", style_name)
            marks = "#" * (int(level.group()) if level else 1)
            text = f"{marks} {text}"
            headings.append(text)
        _emit_line(
            text,
            location_docx(para_index, 0, len(text)),
        )
        para_index += 1

    for table in document.tables:
        table_count += 1
        for row in table.rows:
            pieces: list[str] = []
            cell_spans: list[tuple[int, int, dict[str, Any]]] = []
            pos = 0
            for cell_index, cell in enumerate(row.cells):
                if cell_index:
                    pieces.append(" | ")
                    pos += 3
                cell_text = " ".join(
                    p.text.strip() for p in cell.paragraphs if p.text.strip()
                )
                span_start = pos
                pieces.append(cell_text)
                pos += len(cell_text)
                if cell_text:
                    cell_spans.append(
                        (
                            span_start,
                            pos,
                            location_docx(para_index, span_start, pos),
                        )
                    )
            _emit_line("".join(pieces), None, cell_spans)
            para_index += 1

    raw = "".join(parts).strip("\n")
    return ParsedDocument(
        raw_text=raw,
        format="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        parser=PARSER_PYTHON_DOCX,
        metadata={"paragraph_count": para_index, "table_count": table_count},
        locations=_trim(locations, len(raw)),
    )


# ---------------------------------------------------------------------------
# XLSX -> openpyxl (per-cell locations: {type, sheet, cell})
# ---------------------------------------------------------------------------


def _parse_xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    parts: list[str] = []
    locations: list[CharLocation] = []
    cursor = 0
    sheet_names: list[str] = []
    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            header = f"# Sheet: {sheet.title}"
            parts.append(header)
            cursor += len(header)
            parts.append("\n\n")
            cursor += 2
            for row in sheet.iter_rows():
                cells = [("" if c.value is None else str(c.value)) for c in row]
                if not any(cells):
                    continue
                start = cursor
                pieces: list[str] = []
                cell_spans: list[tuple[int, int, dict[str, Any]]] = []
                pos = 0
                for cell_index, value in enumerate(cells):
                    if cell_index:
                        pieces.append(" | ")
                        pos += 3
                    span_start = pos
                    pieces.append(value)
                    pos += len(value)
                    cell = row[cell_index]
                    if value and cell.value is not None:
                        cell_spans.append(
                            (
                                span_start,
                                pos,
                                location_xlsx(sheet.title, cell.coordinate),
                            )
                        )
                line = "".join(pieces)
                parts.append(line)
                cursor += len(line)
                for span_start, span_end, loc in cell_spans:
                    locations.append(
                        CharLocation(
                            start=start + span_start,
                            end=start + span_end,
                            location=loc,
                        )
                    )
                parts.append("\n")
                cursor += 1
    finally:
        workbook.close()
    raw = "".join(parts).strip("\n")
    return ParsedDocument(
        raw_text=raw,
        format="xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        parser=PARSER_OPENPYXL,
        metadata={"sheets": sheet_names},
        locations=_trim(locations, len(raw)),
    )


# ---------------------------------------------------------------------------
# TXT / MD / CSV -> direct UTF-8 read
# ---------------------------------------------------------------------------


def _parse_text(path: Path) -> ParsedDocument:
    raw_bytes = path.read_bytes()
    try:
        raw = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        raw = raw_bytes.decode("utf-8", errors="replace")
    return ParsedDocument(
        raw_text=raw,
        format="text",
        mime_type=_guess_mime(path),
        parser=PARSER_DIRECT,
        metadata={},
        locations=[
            CharLocation(start=0, end=len(raw), location=location_text(0, len(raw)))
        ],
    )


# ---------------------------------------------------------------------------
# Fallback -> MarkItDown (declared offset limitations)
# ---------------------------------------------------------------------------


def _parse_fallback(path: Path) -> ParsedDocument:
    from markitdown import MarkItDown

    converter = MarkItDown()
    result = converter.convert(str(path))
    raw = result.text_content or ""
    if not raw.strip():
        raise ParserError(f"fallback parser produced no text for {path.name}")
    return ParsedDocument(
        raw_text=raw,
        format="fallback",
        mime_type=_guess_mime(path),
        parser=PARSER_MARKITDOWN,
        fallback_used=True,
        metadata={
            "converter": "markitdown",
            # declared limitation (spec §15): fallback text has no
            # format-level geometry; only text-level char offsets are honest
            "offset_limitations": (
                "markitdown fallback: format locations unavailable, "
                "text char offsets only"
            ),
        },
        locations=[
            CharLocation(start=0, end=len(raw), location=location_text(0, len(raw)))
        ],
    )


def _trim(locations: list[CharLocation], length: int) -> list[CharLocation]:
    trimmed: list[CharLocation] = []
    for entry in locations:
        start = min(max(entry.start, 0), length)
        end = min(max(entry.end, start), length)
        if end > start:
            trimmed.append(CharLocation(start=start, end=end, location=entry.location))
    return trimmed


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def _file_mtime_utc(path: Path) -> datetime:
    try:
        return datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
    except OSError:
        return datetime.now(UTC)


def parse_file(path: str | Path, filename: str | None = None) -> ParsedDocument:
    """Parse any supported file into the common representation.

    Failure strategy (locked): primary parser -> MarkItDown fallback ->
    ParserError. Empty and oversized extractions are rejected with honest
    messages instead of producing a silently useless document.
    """
    file_path = Path(path)
    if not file_path.is_file():
        raise ParserError(f"file not found: {file_path}")
    name = filename or file_path.name
    suffix = Path(name).suffix.lower()

    primary: tuple[str, Any] | None = None
    if suffix in PDF_SUFFIXES:
        primary = ("pdf", _parse_pdf)
    elif suffix in DOCX_SUFFIXES:
        primary = ("docx", _parse_docx)
    elif suffix in XLSX_SUFFIXES:
        primary = ("xlsx", _parse_xlsx)
    elif suffix in TEXT_SUFFIXES:
        primary = ("text", _parse_text)

    if primary is not None:
        label, parser_fn = primary
        try:
            parsed = parser_fn(file_path)
        except Exception as exc:  # noqa: BLE001 - fallback boundary
            # primary parser failed (corrupt / unsupported variant):
            # fall back to MarkItDown, declaring its offset limitations
            try:
                parsed = _parse_fallback(file_path)
                parsed.fallback_used = True
                parsed.metadata["primary_parser_failed"] = label
                parsed.metadata["primary_parser_error"] = str(exc)[:300]
            except Exception as fallback_exc:  # noqa: BLE001
                raise ParserError(
                    f"{name}: could not be parsed as {label} "
                    f"({str(exc)[:200]}) and fallback failed "
                    f"({str(fallback_exc)[:200]})"
                ) from exc
    else:
        try:
            parsed = _parse_fallback(file_path)
        except Exception as exc:  # noqa: BLE001 - fallback boundary
            raise ParserError(f"unsupported or unreadable file {name}: {exc}") from exc

    if parsed.mime_type == "application/octet-stream":
        parsed.mime_type = _guess_mime(file_path)

    if len(parsed.raw_text) > MAX_DOCUMENT_CHARS:
        raise ParserError(
            f"{name}: extracted text too large "
            f"({len(parsed.raw_text)} chars > {MAX_DOCUMENT_CHARS})"
        )
    if not parsed.raw_text.strip():
        if suffix in PDF_SUFFIXES:
            raise ParserError(
                f"{name}: no extractable text (the PDF may be scanned or "
                "image-only; OCR is not part of the standard path)"
            )
        raise ParserError(f"{name}: document contains no extractable text")

    return parsed


def supported_suffixes() -> list[str]:
    return sorted(PDF_SUFFIXES | DOCX_SUFFIXES | XLSX_SUFFIXES | TEXT_SUFFIXES)


def file_provenance(path: str | Path) -> dict[str, Any]:
    """Honest file-level provenance (never invented, never a secret)."""
    file_path = Path(path)
    import hashlib

    digest = hashlib.sha256()
    size = 0
    with file_path.open("rb") as handle:
        while True:
            chunk = handle.read(1 << 20)
            if not chunk:
                break
            size += len(chunk)
            digest.update(chunk)
    return {
        "file_sha256": digest.hexdigest(),
        "file_size": size,
        "file_mtime": _file_mtime_utc(file_path).isoformat(),
    }
